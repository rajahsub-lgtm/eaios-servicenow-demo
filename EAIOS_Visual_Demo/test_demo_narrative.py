"""What the demo asserts on screen must match what the run did.

The headline callout is the most prominent text a panel reads. It was derived
from whether the plan changed, and the branch for "it did not change" assumed
that meant confidence had stayed high — so a 0.45 LOW assessment running the
full investigation was captioned "confidence stayed high, and EAIOS retained
the smallest governed skill plan needed."

A blank panel is a smaller problem than a confident false statement, and no
static check catches this: the app parses, the function returns a string, and
the string is wrong.
"""

from pathlib import Path
import json
import unittest


ROOT = Path(__file__).resolve().parent
BUNDLE = ROOT / "outputs" / "servicenow_story_bundle.json"


def steady_plan_callout():
    """The helper, loaded without executing Streamlit page setup."""
    source = (ROOT / "eaios_story_app.py").read_text(encoding="utf-8")
    start = source.index("def _steady_plan_callout")
    end = source.index("def render_adaptive_story")
    namespace: dict = {"Any": object}
    exec(source[start:end], namespace)
    return namespace["_steady_plan_callout"]


def assessments() -> list[dict]:
    return json.loads(BUNDLE.read_text(encoding="utf-8"))["assessments"]


class TheArcIsCompleteTests(unittest.TestCase):
    def test_every_beat_reaches_the_demo(self):
        """Five scenarios existed in the engine and none of them were in the
        story bundle, so three sprints of work were invisible in the UI."""
        present = {row["scenario_id"] for row in assessments()}
        for scenario_id in (
            "SCN-PAY-001",
            "SCN-PAY-CONTRADICT-001",
            "SCN-PAY-RESOLVED-001",
            "SCN-CROSS-GATEWAY-001",
            "SCN-PAY-EU-001",
            "SCN-NOVEL-INDEX-001",
            "SCN-UNDOCUMENTED-001",
        ):
            with self.subTest(scenario=scenario_id):
                self.assertIn(scenario_id, present)

    def test_the_bundle_opens_on_the_first_beat(self):
        """The selector follows bundle order and the app opens on index 0."""
        self.assertEqual(assessments()[0]["scenario_id"], "SCN-PAY-001")
        source = (ROOT / "eaios_story_app.py").read_text(encoding="utf-8")
        self.assertIn("default_index = 0", source)

    def test_the_recurrence_is_not_a_static_entry(self):
        """It only exists after a human validates and an outcome is recorded.
        A static entry would show it falling back to documentation and
        misrepresent the thing it exists to demonstrate."""
        present = {row["scenario_id"] for row in assessments()}
        self.assertNotIn("SCN-INDEX-RECUR-001", present)


class TheCalloutMatchesTheRunTests(unittest.TestCase):
    def test_no_scenario_is_told_confidence_stayed_high_when_it_did_not(self):
        callout = steady_plan_callout()
        for row in assessments():
            if row["expanded_during_execution"] or row["contracted_during_execution"]:
                continue
            with self.subTest(scenario=row["correlation_id"]):
                text = callout(row)
                if row["final_confidence_level"] != "HIGH":
                    self.assertNotIn("confidence stayed high", text)
                if row["final_plan_mode"] != "ACCELERATED_VALIDATION":
                    self.assertNotIn("smallest governed skill plan", text)

    def test_each_steady_scenario_gets_its_own_claim(self):
        callout = steady_plan_callout()
        expected = {
            "EAIOS-DEMO-STABLE-001": "Efficient behavior",
            "EAIOS-DEMO-TRANSFERRED-001": "Transferred experience",
            "EAIOS-DEMO-UNSEEN-001": "Reasoning from documentation",
            "EAIOS-DEMO-UNDOCUMENTED-001": "Governed non-diagnosis",
        }
        by_id = {row["correlation_id"]: row for row in assessments()}
        for correlation_id, claim in expected.items():
            with self.subTest(scenario=correlation_id):
                self.assertIn(claim, callout(by_id[correlation_id]))


class TheMarkupDoesNotLeakTests(unittest.TestCase):
    def test_the_plan_revision_html_has_no_indented_lines(self):
        """Indented HTML inside st.markdown becomes a code block and renders
        as literal text. It only showed when the cancelled block was empty,
        which is most scenarios, including the one the demo opens on."""
        source = (ROOT / "eaios_story_app.py").read_text(encoding="utf-8")
        start = source.index('<div class="eaios-section-label">Adaptive plan revision</div>')
        block = source[start : source.index('"""', start)]
        for line in block.splitlines()[1:]:
            with self.subTest(line=line[:50]):
                self.assertFalse(
                    line.startswith("    "),
                    "indented line inside an HTML markdown block",
                )


class DependenciesAreInstallableTests(unittest.TestCase):
    def test_every_app_import_is_declared(self):
        """plotly was required by the app and absent from the environment, so
        the demo could not start at all."""
        declared = (ROOT / "requirements.txt").read_text(encoding="utf-8")
        for package in ("streamlit", "plotly", "pandas"):
            with self.subTest(package=package):
                self.assertIn(package, declared)

    def test_they_are_actually_importable(self):
        for module in ("streamlit", "plotly.express", "pandas"):
            with self.subTest(module=module):
                __import__(module)


if __name__ == "__main__":
    unittest.main()


class TheScriptQuotesRealNumbersTests(unittest.TestCase):
    """A talk track with stale numbers is worse than no talk track.

    The presenter reads them aloud with the screen beside them, so a figure
    that has drifted is contradicted in front of the audience by the demo
    itself.
    """

    SCRIPT = ROOT / "DEMO_SCRIPT.md"

    def script(self) -> str:
        return self.SCRIPT.read_text(encoding="utf-8")

    def test_every_quoted_confidence_appears_in_the_bundle(self):
        import re

        quoted = set(re.findall(r"`?(\d\.\d{3})", self.script()))
        actual = set()
        for row in assessments():
            actual.add(f"{row['initial_confidence_score']:.3f}")
            actual.add(f"{row['final_confidence_score']:.3f}")
        # 0.341 and 0.000 come from the live learning loop, not the bundle.
        live = {"0.341", "0.000", "0.220", "0.100", "0.950", "0.450"}
        for number in quoted:
            with self.subTest(number=number):
                self.assertTrue(
                    number in actual or number in live,
                    f"{number} is quoted in the script and matches nothing",
                )

    def test_the_beats_named_in_the_script_are_the_beats_in_the_bundle(self):
        script = self.script()
        for label, _, _ in [
            ("Payment connector — stable evidence", None, None),
            ("Payment connector (EU) — experience by analogy", None, None),
            ("Search indexer — nothing comparable", None, None),
            ("Licence reconciler — nothing written either", None, None),
        ]:
            with self.subTest(label=label):
                self.assertIn(label, script)

    def test_the_rehearsal_command_exists(self):
        self.assertIn("rehearse.py", self.script())
        self.assertTrue((ROOT / "rehearse.py").exists())

    def test_the_script_does_not_promise_autonomous_action(self):
        """Every path requires approval; the script must not imply otherwise."""
        script = self.script().lower()
        for phrase in ("automatically remediat", "without approval", "acts on its own"):
            with self.subTest(phrase=phrase):
                self.assertNotIn(phrase, script)


class TheDocumentationMatchesTheSystemTests(unittest.TestCase):
    """Figures typed into prose drift. Figures in the deck are read from the
    fixtures; the ones in the architecture note were typed, and were wrong on
    first writing — 63 entities against an actual 30.
    """

    DOC = ROOT / "ARCHITECTURE.md"

    def counts(self) -> dict:
        return {
            name: len(
                json.loads(
                    (ROOT / "json" / f"{name}.json").read_text(encoding="utf-8")
                )
            )
            for name in (
                "entities",
                "semantic_relationships",
                "outcome_history",
                "knowledge_documents",
                "scenarios",
            )
        }

    def test_the_quoted_scale_is_the_real_scale(self):
        import re

        # Normalised, because the figure and its label are routinely split
        # across a line break and the prose is wrapped for reading.
        doc = re.sub(r"\s+", " ", self.DOC.read_text(encoding="utf-8"))
        counts = self.counts()
        for label, key in (
            ("entities", "entities"),
            ("governed relationships", "semantic_relationships"),
            ("recorded outcomes", "outcome_history"),
            ("knowledge documents", "knowledge_documents"),
            ("scenarios", "scenarios"),
        ):
            with self.subTest(label=label):
                self.assertIn(f"{counts[key]} {label}", doc)

    def test_the_deck_builder_reads_rather_than_states(self):
        """The deck must fail to build rather than misrepresent the system."""
        source = (ROOT / "build_deck.py").read_text(encoding="utf-8")
        self.assertIn("def facts()", source)
        self.assertIn('F[\'entities\']', source)

    def test_the_deck_exists_and_has_every_section(self):
        from pptx import Presentation

        deck = ROOT / "outputs" / "EAIOS_Architecture_and_Demo.pptx"
        self.assertTrue(deck.exists(), "run: python build_deck.py")
        text = " ".join(
            shape.text_frame.text
            for slide in Presentation(deck).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        for concept in (
            "OPERATIONAL CONFIDENCE",
            "DYNAMIC PLANNING",
            "COLLECTIVE INTELLIGENCE",
            "GOVERNANCE",
            "EVIDENCE FUSION",
            "KNOWLEDGE GRAPH",
            "LEARNING FROM OUTCOME",
        ):
            with self.subTest(concept=concept):
                self.assertIn(concept, text)


class TheDesignDocumentIsAccurateTests(unittest.TestCase):
    """Design documents rot silently. This one quotes forty-odd parameters
    that live in policy, and a reader has no way to tell a stale figure from
    a current one.
    """

    DESIGN = ROOT / "DESIGN.md"

    def doc(self) -> str:
        return self.DESIGN.read_text(encoding="utf-8")

    def policy(self, name: str) -> dict:
        return json.loads(
            (ROOT / "config" / f"{name}.json").read_text(encoding="utf-8")
        )

    def test_the_confidence_weights_are_current(self):
        doc = self.doc()
        for name, weight in self.policy("confidence_policy")["weights"].items():
            with self.subTest(factor=name):
                self.assertIn(f"`{name}` | {weight:.2f}", doc)

    def test_the_similarity_weights_are_current(self):
        doc = self.doc()
        weights = self.policy("similarity_policy")["dimension_weights"]
        for name, weight in weights.items():
            with self.subTest(dimension=name):
                self.assertIn(f"`{name}` | {weight:.2f}", doc)

    def test_every_blocking_flag_is_documented(self):
        """A flag that bars the narrow plan and is absent from the governance
        section is a rule nobody reading this could discover."""
        doc = self.doc()
        accelerated = next(
            mode
            for mode in self.policy("orchestration_policies")["modes"]
            if mode["mode_id"] == "ACCELERATED_VALIDATION"
        )
        for flag in accelerated["entry_conditions"]["disallowed_hard_flags"]:
            with self.subTest(flag=flag):
                self.assertIn(flag, doc)

    def test_the_regression_baseline_matches_the_bundle(self):
        doc = self.doc()
        for row in assessments():
            with self.subTest(scenario=row["scenario_id"]):
                if row["scenario_id"] not in doc:
                    continue
                self.assertIn(
                    f"{row['initial_confidence_score']:.3f} → "
                    f"{row['final_confidence_score']:.3f}",
                    doc,
                )

    def test_the_named_modules_exist(self):
        import re

        doc = self.doc()
        for module in sorted(set(re.findall(r"`([a-z_]+\.py)`", doc))):
            with self.subTest(module=module):
                self.assertTrue(
                    (ROOT / module).exists(), f"{module} is documented but absent"
                )

    # Written by a run rather than shipped, so they are absent until the
    # learning loop creates them. Section 4.3 documents them as exactly that.
    RUNTIME_WRITTEN = {
        "runtime_outcome_feedback.json",
        "learned_patterns.json",
        "refutation_ledger.json",
    }

    def test_the_named_policy_files_exist(self):
        import re

        doc = self.doc()
        named = set(re.findall(r"`([a-z_]+\.json)`", doc)) - self.RUNTIME_WRITTEN
        for name in sorted(named):
            with self.subTest(policy=name):
                self.assertTrue(
                    (ROOT / "config" / name).exists()
                    or (ROOT / "json" / name).exists(),
                    f"{name} is documented but absent",
                )


class TheApprovalBoundaryHoldsWithoutAnInstanceTests(unittest.TestCase):
    """Human approval is the architecture, not a ServiceNow feature.

    The demonstration used to deep-link into a personal developer instance.
    Removing it must not weaken the boundary — if approval only held because a
    particular tenant was configured, it was never a property of the system.
    """

    def test_no_tenant_specific_reference_ships(self):
        """A developer instance name and record sys_ids are tenant identifiers
        and have no place in a deployed image."""
        import re

        for path in list(ROOT.glob("*.py")) + list((ROOT / "config").glob("*.json")):
            text = path.read_text(encoding="utf-8")
            with self.subTest(file=path.name):
                self.assertNotRegex(
                    text,
                    r"dev\d{6}\.service-now\.com",
                    "a specific developer instance is referenced",
                )

    def test_the_record_panel_survives_an_unbound_instance(self):
        config = json.loads(
            (ROOT / "config" / "visual_demo.json").read_text(encoding="utf-8")
        )
        self.assertEqual(config["instance_url"], "")
        self.assertEqual(config["records"], {})

        from eaios_story_data import StoryRepository

        repo = StoryRepository.load(ROOT)
        for row in assessments():
            with self.subTest(scenario=row["correlation_id"]):
                # No link, and no exception — the panel falls back to the
                # contract rather than failing.
                self.assertIsNone(
                    repo.servicenow_record_url(row["correlation_id"])
                )
                self.assertEqual(
                    repo.servicenow_record_metadata(row["correlation_id"]), {}
                )

    def test_every_path_still_requires_a_human(self):
        """The boundary is unconditional. Confidence, recall class and plan
        width change what is investigated, never whether approval is needed."""
        for row in assessments():
            with self.subTest(scenario=row["correlation_id"]):
                self.assertTrue(
                    row["recommendation"]["human_approval_required"],
                    "a path reached a recommendation without requiring approval",
                )
                self.assertEqual(
                    row["safety_status"], "REQUIRES_HUMAN_APPROVAL"
                )

    def test_the_write_is_escalated_rather_than_performed(self):
        for row in assessments():
            escalations = [
                d
                for d in row.get("policy_decisions", [])
                if d.get("decision") == "ESCALATE"
            ]
            with self.subTest(scenario=row["correlation_id"]):
                self.assertTrue(
                    escalations,
                    "creating the assessment record was not escalated",
                )
